from pathlib import Path
from openpyxl import load_workbook

ARQUIVO = Path(__file__).parent / "matriz_scoring_procurement.xlsx"

WEIGHTS = {
    "Nota Preço": 0.30,
    "Nota SLA": 0.25,
    "Nota Pagamento": 0.15,
    "Nota Escopo": 0.15,
    "Nota Garantia": 0.10,
    "Nota Cobertura": 0.05,
}


def ler_aba(wb, nome_aba, linha_cabecalho=1):
    """Lê uma aba e devolve lista de dicts {coluna: valor}."""
    rows = list(wb[nome_aba].iter_rows(values_only=True))
    headers = [str(h).strip() if h else "" for h in rows[linha_cabecalho - 1]]
    dados = []
    for r in rows[linha_cabecalho:]:
        if r and r[0]:
            dados.append(dict(zip(headers, r)))
    return dados


def calcular_score(v):
    return sum(v[coluna] * peso for coluna, peso in WEIGHTS.items())


def carregar_dados(caminho):
    # Sem data_only: a coluna "Score Final" é fórmula sem valor salvo,
    # então o score é calculado aqui em Python.
    wb = load_workbook(caminho)

    scoring = ler_aba(wb, "Scoring", linha_cabecalho=4)
    # Só linhas de fornecedor (preço numérico), ignora o bloco "Regras de pontuação"
    scoring = [v for v in scoring if isinstance(v.get("Preço (R$)"), (int, float))]

    # Nota Preço automática: menor preço / preço do fornecedor x 100
    menor_preco = min(v["Preço (R$)"] for v in scoring)
    for v in scoring:
        v["Nota Preço"] = menor_preco / v["Preço (R$)"] * 100
    compliance = {c["Fornecedor"]: c for c in ler_aba(wb, "Compliance")}
    dados = {d["Fornecedor"]: d for d in ler_aba(wb, "Dados_e_Notas")}
    resultado = {r["Fornecedor"]: r for r in ler_aba(wb, "Resultado", linha_cabecalho=3)}

    analises = []
    for v in scoring:
        nome = v["Fornecedor"]
        analises.append({
            "fornecedor": nome,
            "preco": v["Preço (R$)"],
            "score": round(calcular_score(v), 2),
            "compliance": compliance[nome]["Status"],
            "leitura": resultado[nome]["Leitura"],
            "oportunidade": resultado[nome]["Oportunidade"],
            "ressalva": resultado[nome]["Ressalva"],
            "escopo_texto": dados[nome]["Escopo"],
            "nota_escopo": v["Nota Escopo"],
        })
    return analises


def gerar_relatorio(analises):
    # Conformes primeiro, depois por score
    analises = sorted(
        analises,
        key=lambda a: (a["compliance"] != "CONFORME", -a["score"]),
    )

    linhas = [
        "ANÁLISE AUTOMÁTICA — Supplier Scoring & Ranking Tool",
        "O score apoia a decisão e não substitui validação técnica/comercial.",
        "",
    ]
    for i, a in enumerate(analises, 1):
        linhas.append(
            f"{i}. {a['fornecedor']} | Score {a['score']} | "
            f"R$ {a['preco']:,.0f} | {a['compliance']}".replace(",", ".")
        )
        linhas.append(f"   Leitura: {a['leitura']}")
        linhas.append(f"   Oportunidade: {a['oportunidade']}")
        linhas.append(f"   Ressalva: {a['ressalva']}")

    # Alerta de inconsistência: escopo com ressalva mas nota 100
    alertas = [
        a for a in analises
        if str(a["escopo_texto"]).strip().lower() != "atende" and a["nota_escopo"] == 100
    ]
    if alertas:
        linhas.append("")
        linhas.append("ATENÇÃO — possível inconsistência na planilha:")
        for a in alertas:
            linhas.append(
                f"   {a['fornecedor']}: escopo = '{a['escopo_texto']}' "
                f"mas Nota Escopo = {a['nota_escopo']}"
            )

    return "\n".join(linhas)


if __name__ == "__main__":
    analises = carregar_dados(ARQUIVO)
    print(gerar_relatorio(analises))
